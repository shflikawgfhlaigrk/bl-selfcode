"""ops/export_leads_xlsx.py — the leads XLSX exporter's DB boundary is bounded,
injectable and never raises; main() exits honestly on every failure. No live
Postgres and no openpyxl required (workbook tests skip if openpyxl is absent)."""
from __future__ import annotations

import importlib.util
import pathlib
import subprocess
from types import SimpleNamespace

import psycopg
import pytest

from utah import config

_PATH = pathlib.Path(__file__).resolve().parents[1] / "ops" / "export_leads_xlsx.py"
_spec = importlib.util.spec_from_file_location("ops_export_leads_xlsx", _PATH)
exporter = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(exporter)


# --- fakes ------------------------------------------------------------------

_COLS = ("id", "name", "kind", "region", "source", "status",
         "phone", "email", "website", "address", "ts")


class _FakeCursor:
    description = [SimpleNamespace(name=n) for n in _COLS]

    def __init__(self, rows):
        self._rows = rows

    def fetchall(self):
        return self._rows


class _FakeConn:
    def __init__(self, rows):
        self._rows = rows

    def __enter__(self):
        return self

    def __exit__(self, *a):
        return False

    def execute(self, sql):
        return _FakeCursor(self._rows)


def _tuple_row():
    return (1, "Acme Plumbing", "plumber", "Augusta GA", "osm", "new",
            "+15551234567", "a@b.c", None, "1 Main St", "2026-06-01 00:00:00+00")


def _dictrow(**over):
    d = {"id": "1", "name": "Acme Plumbing", "kind": "plumber",
         "region": "Augusta GA", "source": "osm", "status": "new",
         "phone": "+15551234567", "email": "a@b.c", "website": "",
         "address": "1 Main St", "ts": "2026-06-01 00:00:00+00"}
    d.update(over)
    return d


def _ok_fetch(rows):
    return lambda: {"ok": True, "rows": rows, "via": "test", "error": None}


# --- psycopg lane -------------------------------------------------------------

def test_fetch_rows_psycopg_bounded_connect_and_normalized_rows():
    seen = {}

    def connect(dsn, connect_timeout=None, options=None):
        seen["dsn"] = dsn
        seen["timeout"] = connect_timeout
        seen["options"] = options
        return _FakeConn([_tuple_row()])

    res = exporter.fetch_rows(connect=connect)
    assert res["ok"] is True and res["via"] == "psycopg" and res["error"] is None
    assert seen["dsn"] == config.DB_DSN
    assert seen["timeout"] == config.DB_CONNECT_TIMEOUT   # bounded — never hangs
    assert "statement_timeout" in (seen["options"] or ""), \
        "the full-table SELECT must be statement-bounded, not just the connect"
    row = res["rows"][0]
    assert row["name"] == "Acme Plumbing"
    assert row["website"] == ""    # NULL -> "" exactly like the psql CSV lane
    assert row["id"] == "1"        # one stringly shape for both lanes


def test_fetch_rows_db_failure_is_honest_not_raised():
    def connect(dsn, connect_timeout=None, options=None):
        raise psycopg.OperationalError("connection refused")

    res = exporter.fetch_rows(connect=connect)
    assert res["ok"] is False and res["rows"] == []
    assert "connection refused" in res["error"]


# --- psql fallback lane -------------------------------------------------------

def test_psql_lane_is_bounded_and_parses_csv():
    calls = {}
    csv_text = ",".join(_COLS) + "\n" + "2,Bravo Roofing,roofer,Macon GA,osm,new,,,,,\n"

    def runner(cmd, **kw):
        calls["cmd"] = cmd
        calls["kw"] = kw
        return SimpleNamespace(stdout=csv_text, returncode=0)

    res = exporter._fetch_via_psql(runner)
    assert res["ok"] is True and res["via"] == "psql"
    assert res["rows"][0]["name"] == "Bravo Roofing"
    kw = calls["kw"]
    assert kw["timeout"] > 0                              # bounded subprocess
    assert kw["env"]["PGCONNECT_TIMEOUT"]                 # bounded connect
    assert "ON_ERROR_STOP=1" in calls["cmd"]              # honest non-zero on SQL error


def test_psql_lane_failure_is_honest_not_raised():
    def runner(cmd, **kw):
        raise subprocess.CalledProcessError(
            2, cmd, stderr="psql: error: connection to server failed")

    res = exporter._fetch_via_psql(runner)
    assert res["ok"] is False and res["rows"] == []
    assert "connection to server failed" in res["error"]


def test_psql_lane_timeout_is_honest_not_raised():
    def runner(cmd, **kw):
        raise subprocess.TimeoutExpired(cmd, 120)

    res = exporter._fetch_via_psql(runner)
    assert res["ok"] is False
    assert "120" in res["error"]


def test_fetch_rows_routes_to_psql_when_psycopg_missing(monkeypatch):
    monkeypatch.setattr(exporter, "psycopg", None)
    csv_text = ",".join(_COLS) + "\n" + "3,Charlie HVAC,hvac,Albany GA,osm,new,,,,,\n"

    def runner(cmd, **kw):
        return SimpleNamespace(stdout=csv_text, returncode=0)

    res = exporter.fetch_rows(runner=runner)
    assert res["ok"] is True and res["via"] == "psql"
    assert res["rows"][0]["name"] == "Charlie HVAC"


# --- pure aggregation ---------------------------------------------------------

def test_summarize_regions_counts_and_orders_desc():
    rows = [{"region": "A", "email": "x@y", "phone": ""},
            {"region": "A", "email": "", "phone": "1"},
            {"region": "B", "email": "", "phone": ""}]
    out = exporter.summarize_regions(rows)
    assert out == [("A", 2, 1, 1), ("B", 1, 0, 0)]


# --- main() honesty -----------------------------------------------------------

def test_main_exits_nonzero_and_reports_to_stderr_on_db_failure(tmp_path, capsys):
    out = tmp_path / "x.xlsx"
    rc = exporter.main(
        [str(out)],
        fetch=lambda: {"ok": False, "rows": [], "via": "psycopg", "error": "db down"})
    assert rc != 0
    assert "db down" in capsys.readouterr().err
    assert not out.exists()


def test_main_honest_when_openpyxl_missing(tmp_path, capsys, monkeypatch):
    monkeypatch.setattr(exporter, "Workbook", None)
    rc = exporter.main([str(tmp_path / "x.xlsx")], fetch=_ok_fetch([_dictrow()]))
    assert rc != 0
    assert "openpyxl" in capsys.readouterr().err


def test_main_happy_path_writes_workbook(tmp_path, capsys):
    pytest.importorskip("openpyxl")
    from openpyxl import load_workbook

    out = tmp_path / "leads.xlsx"
    rows = [_dictrow(), _dictrow(id="2", name="Bravo", email="", status="worked")]
    rc = exporter.main([str(out)], fetch=_ok_fetch(rows))
    assert rc == 0 and out.exists()
    assert "Exported 2 leads" in capsys.readouterr().out
    wb = load_workbook(out)
    assert wb.sheetnames == ["All Leads", "With Email", "With Phone",
                             "Never Contacted", "By Region"]
    assert wb["All Leads"].max_row == 3        # header + 2 rows
    assert wb["With Email"].max_row == 2       # only the lead with an email
    assert wb["Never Contacted"].max_row == 2  # only status=new


def test_main_save_failure_is_honest(tmp_path, capsys):
    pytest.importorskip("openpyxl")
    out = tmp_path / "missing-dir" / "x.xlsx"
    rc = exporter.main([str(out)], fetch=_ok_fetch([_dictrow()]))
    assert rc != 0
    assert "x.xlsx" in capsys.readouterr().err
