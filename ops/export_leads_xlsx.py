#!/usr/bin/env python3
"""Mass-export Utah leads to an Excel workbook: All Leads + ready-to-work sheets
(With Email, With Phone, Never Contacted) + a per-region summary.

Usage: python3 ops/export_leads_xlsx.py [output.xlsx]
Default output: ~/Desktop/leads-export-YYYY-MM-DD.xlsx

Boundaries are bounded and honest (RUBRIC hard caps): the DB read prefers a
native psycopg connect with ``connect_timeout=config.DB_CONNECT_TIMEOUT``
(the utah/product/tasks.py pattern) and falls back to a psql subprocess that
gets BOTH ``PGCONNECT_TIMEOUT`` and a subprocess timeout — a stalled cluster
can never hang the export. Every failure prints to stderr and exits non-zero;
nothing is swallowed by a bare except.
"""
from __future__ import annotations

import csv
import datetime
import io
import os
import pathlib
import subprocess
import sys

# The repo root rides on sys.path so `utah.config` resolves when this script is
# run directly (path is derived from __file__ — no hardcoded machine paths).
_REPO_ROOT = pathlib.Path(__file__).resolve().parents[1]
if str(_REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(_REPO_ROOT))

from utah import config  # noqa: E402  (needs the sys.path seam above)

# Optional lanes, probed honestly: the repo venv has psycopg but no openpyxl;
# the system python3 has openpyxl but no psycopg. Each side gates explicitly.
try:
    import psycopg
except ModuleNotFoundError:  # → bounded psql subprocess lane
    psycopg = None

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
except ModuleNotFoundError:  # main() reports honestly instead of crashing
    Workbook = None

HEADERS = ["id", "name", "kind", "region", "source", "status",
           "phone", "email", "website", "address", "ts"]

_SELECT = """SELECT id, name, kind, region, source, status,
       contact->>'phone'   AS phone,
       contact->>'email'   AS email,
       contact->>'website' AS website,
       contact->>'address' AS address,
       ts
FROM leads ORDER BY region, name"""

_PSQL_COPY = "\\copy (" + _SELECT + ") TO STDOUT WITH (FORMAT csv, HEADER true)"
_PSQL_GLOB = "opt/homebrew/opt/postgresql@*/bin/psql"
#: Hard wall for the whole psql subprocess (connect is separately bounded by
#: PGCONNECT_TIMEOUT). Env-overridable for huge tables.
_PSQL_TIMEOUT_S = int(os.environ.get("UTAH_EXPORT_PSQL_TIMEOUT", "120"))


def psql_bin() -> str:
    hits = sorted(pathlib.Path("/").glob(_PSQL_GLOB))
    return str(hits[-1]) if hits else "psql"


def _normalize(names: list[str], raw: list[tuple]) -> list[dict]:
    """One stringly row shape for both lanes (NULL -> '' like psql CSV)."""
    return [{n: ("" if v is None else str(v)) for n, v in zip(names, row)}
            for row in raw]


def _fetch_via_psycopg(connect) -> dict:
    """Bounded native read. Never raises — {"ok": bool, "rows": [...], ...}."""
    errors = (psycopg.Error, OSError) if psycopg is not None else (OSError,)
    try:
        with connect(config.DB_DSN,
                     connect_timeout=config.DB_CONNECT_TIMEOUT,
                     options=f"-c statement_timeout={config.DB_STATEMENT_TIMEOUT_MS}") as conn:
            cur = conn.execute(_SELECT)
            names = [c.name for c in cur.description]
            raw = cur.fetchall()
    except errors as exc:
        return {"ok": False, "rows": [], "via": "psycopg",
                "error": f"leads query failed ({config.DB_DSN}): {exc}"}
    return {"ok": True, "rows": _normalize(names, raw), "via": "psycopg",
            "error": None}


def _fetch_via_psql(runner) -> dict:
    """Bounded psql subprocess read. Never raises — same shape as above."""
    env = dict(os.environ)
    env.setdefault("PGCONNECT_TIMEOUT", str(config.DB_CONNECT_TIMEOUT))
    cmd = [psql_bin(), "-h", "127.0.0.1", "-p", "5433", "-d", "utah",
           "-v", "ON_ERROR_STOP=1", "-c", _PSQL_COPY]
    try:
        out = runner(cmd, capture_output=True, text=True, check=True,
                     timeout=_PSQL_TIMEOUT_S, env=env)
    except subprocess.TimeoutExpired as exc:
        return {"ok": False, "rows": [], "via": "psql",
                "error": f"psql export exceeded {exc.timeout}s — aborted"}
    except (subprocess.CalledProcessError, OSError) as exc:
        stderr = (getattr(exc, "stderr", "") or "").strip()
        return {"ok": False, "rows": [], "via": "psql",
                "error": f"psql export failed: {stderr or exc}"}
    return {"ok": True, "rows": list(csv.DictReader(io.StringIO(out.stdout))),
            "via": "psql", "error": None}


def fetch_rows(*, connect=None, runner=None) -> dict:
    """Read the leads table. Never raises; prefers psycopg, falls back to psql.

    ``connect``/``runner`` are injectable test seams (default: real psycopg
    connect / subprocess.run).
    """
    if connect is not None or psycopg is not None:
        return _fetch_via_psycopg(connect or psycopg.connect)
    return _fetch_via_psql(runner or subprocess.run)


def summarize_regions(rows: list[dict]) -> list[tuple]:
    """(region, leads, with_email, with_phone) tuples, biggest region first."""
    regions: dict[str, list[int]] = {}
    for r in rows:
        agg = regions.setdefault(r["region"], [0, 0, 0])
        agg[0] += 1
        agg[1] += bool(r["email"])
        agg[2] += bool(r["phone"])
    return [(region, n, e, p) for region, (n, e, p) in
            sorted(regions.items(), key=lambda kv: -kv[1][0])]


def add_sheet(wb, title: str, rows: list[dict]) -> None:
    ws = wb.create_sheet(title)
    ws.append(HEADERS)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append([r.get(h, "") for h in HEADERS])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{len(rows) + 1}"
    widths = {"name": 34, "kind": 16, "region": 30, "phone": 15, "email": 30,
              "website": 34, "address": 40, "ts": 22}
    for i, h in enumerate(HEADERS, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(h, 10)


def build_workbook(rows: list[dict]):
    wb = Workbook()
    wb.remove(wb.active)
    add_sheet(wb, "All Leads", rows)
    add_sheet(wb, "With Email", [r for r in rows if r["email"]])
    add_sheet(wb, "With Phone", [r for r in rows if r["phone"]])
    add_sheet(wb, "Never Contacted", [r for r in rows if r["status"] == "new"])

    summary = wb.create_sheet("By Region")
    summary.append(["region", "leads", "with_email", "with_phone"])
    for c in summary[1]:
        c.font = Font(bold=True)
    for region, n, e, p in summarize_regions(rows):
        summary.append([region, n, e, p])
    summary.column_dimensions["A"].width = 34
    return wb


def main(argv: list[str] | None = None, *, fetch=None) -> int:
    args = list(sys.argv[1:] if argv is None else argv)
    out = pathlib.Path(args[0]) if args else (
        pathlib.Path.home() / "Desktop"
        / f"leads-export-{datetime.date.today()}.xlsx")

    fetched = (fetch or fetch_rows)()
    if not fetched["ok"]:
        print(f"export failed [{fetched['via']}]: {fetched['error']}",
              file=sys.stderr)
        return 1
    if Workbook is None:
        print("export failed: openpyxl is not installed for "
              f"{sys.executable} — run with an interpreter that has it",
              file=sys.stderr)
        return 2

    rows = fetched["rows"]
    try:
        build_workbook(rows).save(out)
    except OSError as exc:
        print(f"export failed: could not write {out}: {exc}", file=sys.stderr)
        return 1
    print(f"Exported {len(rows)} leads -> {out} "
          f"(email {sum(1 for r in rows if r['email'])}, "
          f"phone {sum(1 for r in rows if r['phone'])})")
    return 0


if __name__ == "__main__":
    sys.exit(main())
